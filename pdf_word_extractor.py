#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
智能PDF单词提取器
自动从PDF词汇书中提取单词、音标、释义、例句等信息
"""

import re
import sys
from pathlib import Path

try:
    import pdfplumber
except ImportError:
    print("正在安装必要的库 pdfplumber...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "pdfplumber", "openpyxl"])
    import pdfplumber

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
except ImportError:
    print("正在安装 openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill


class WordEntry:
    """单词条目类"""
    def __init__(self):
        self.word = ""
        self.phonetic = ""
        self.definition = ""
        self.example_en = ""
        self.example_cn = ""
        self.tags = []

    def __repr__(self):
        return f"Word({self.word}, {self.phonetic}, {self.definition[:20]}...)"


class PDFWordExtractor:
    """PDF单词提取器"""

    def __init__(self, pdf_path):
        self.pdf_path = pdf_path
        self.words = []

    def extract_text(self):
        """提取PDF文本"""
        print(f"正在读取PDF: {self.pdf_path}")
        all_text = []

        with pdfplumber.open(self.pdf_path) as pdf:
            total_pages = len(pdf.pages)
            print(f"PDF共有 {total_pages} 页")

            for i, page in enumerate(pdf.pages, 1):
                if i % 10 == 0:
                    print(f"处理进度: {i}/{total_pages} 页")

                text = page.extract_text()
                if text:
                    all_text.append(text)

        return '\n'.join(all_text)

    def parse_words(self, text):
        """智能解析单词"""
        print("\n正在智能解析单词...")

        # 按行分割
        lines = text.split('\n')

        current_entry = None

        for line in lines:
            line = line.strip()
            if not line:
                continue

            # 检测是否是新单词（以英文字母开头，后面可能跟音标）
            word_pattern = r'^([a-zA-Z][a-zA-Z\s\-\']+)\s*[\[/]?([^\]\/\n]*?)[\]\/]?'
            match = re.match(word_pattern, line)

            if match and self._is_likely_word(match.group(1)):
                # 保存上一个单词
                if current_entry and current_entry.word:
                    self.words.append(current_entry)

                # 开始新单词
                current_entry = WordEntry()
                current_entry.word = match.group(1).strip()

                # 尝试提取音标
                phonetic = self._extract_phonetic(line)
                if phonetic:
                    current_entry.phonetic = phonetic

                continue

            # 如果当前有活跃的单词条目
            if current_entry:
                # 检测音标
                if not current_entry.phonetic:
                    phonetic = self._extract_phonetic(line)
                    if phonetic:
                        current_entry.phonetic = phonetic
                        continue

                # 检测中文释义
                if re.search(r'[\u4e00-\u9fa5]', line):
                    # 如果是英文例句的翻译
                    if current_entry.example_en and not current_entry.example_cn:
                        current_entry.example_cn = line
                    # 如果看起来是释义（短句子，有中文）
                    elif len(line) < 100 and not current_entry.definition:
                        # 清理词性标记等
                        definition = re.sub(r'^[a-z]+\.?\s*', '', line)
                        current_entry.definition = definition.strip()
                    # 可能是额外的释义
                    elif current_entry.definition:
                        current_entry.definition += "; " + line

                # 检测英文例句
                elif re.match(r'^[A-Z].*[.!?]$', line) and len(line) > 10:
                    if not current_entry.example_en:
                        current_entry.example_en = line

        # 保存最后一个单词
        if current_entry and current_entry.word:
            self.words.append(current_entry)

        print(f"成功提取 {len(self.words)} 个单词！")
        return self.words

    def _is_likely_word(self, text):
        """判断是否可能是单词"""
        text = text.strip().lower()

        # 排除常见的非单词模式
        exclude_patterns = [
            r'^\d+$',  # 纯数字
            r'^page\s+\d+',  # 页码
            r'^chapter',  # 章节
            r'^unit',  # 单元
            r'^part\s+[ivxlcdm]+',  # 部分
        ]

        for pattern in exclude_patterns:
            if re.match(pattern, text):
                return False

        # 单词应该不太长
        if len(text) > 30:
            return False

        # 应该主要由字母组成
        alpha_count = sum(c.isalpha() for c in text)
        if alpha_count < len(text) * 0.7:
            return False

        return True

    def _extract_phonetic(self, text):
        """提取音标"""
        # 匹配 /.../ 或 [...] 或 【...】格式的音标
        patterns = [
            r'/([^/]+)/',
            r'\[([^\]]+)\]',
            r'【([^】]+)】',
        ]

        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                return '/' + match.group(1).strip() + '/'

        return ""

    def export_to_excel(self, output_path, include_fields=None):
        """导出为Excel"""
        if include_fields is None:
            include_fields = ['word', 'phonetic', 'definition', 'example_en', 'example_cn']

        print(f"\n正在导出到Excel: {output_path}")

        wb = Workbook()
        ws = wb.active
        ws.title = "单词表"

        # 字段映射
        field_names = {
            'word': '单词',
            'phonetic': '音标',
            'definition': '释义',
            'example_en': '英文例句',
            'example_cn': '中文例句',
        }

        # 写入表头
        headers = [field_names[field] for field in include_fields]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, size=12)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(bold=True, size=12, color="FFFFFF")
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # 写入数据
        for row_idx, word_entry in enumerate(self.words, 2):
            for col_idx, field in enumerate(include_fields, 1):
                value = getattr(word_entry, field, '')
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(vertical='top', wrap_text=True)

        # 调整列宽
        column_widths = {
            'word': 20,
            'phonetic': 25,
            'definition': 40,
            'example_en': 50,
            'example_cn': 50,
        }

        for col_idx, field in enumerate(include_fields, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = column_widths[field]

        wb.save(output_path)
        print(f"✅ 导出成功！文件保存在: {output_path}")
        print(f"共导出 {len(self.words)} 个单词")

    def export_to_txt(self, output_path):
        """导出为TXT"""
        print(f"\n正在导出到TXT: {output_path}")

        with open(output_path, 'w', encoding='utf-8') as f:
            for word in self.words:
                f.write(f"{word.word}")
                if word.phonetic:
                    f.write(f" {word.phonetic}")
                if word.definition:
                    f.write(f" {word.definition}")
                f.write("\n")

                if word.example_en:
                    f.write(f"  {word.example_en}\n")
                if word.example_cn:
                    f.write(f"  {word.example_cn}\n")

                f.write("\n")

        print(f"✅ 导出成功！文件保存在: {output_path}")


def main():
    print("=" * 60)
    print("📚 智能PDF单词提取器")
    print("=" * 60)
    print()

    # 获取PDF文件路径
    if len(sys.argv) > 1:
        pdf_path = sys.argv[1]
    else:
        pdf_path = input("请输入PDF文件路径（拖拽文件到这里也可以）: ").strip().strip('"')

    if not Path(pdf_path).exists():
        print(f"❌ 错误：找不到文件 {pdf_path}")
        input("按回车键退出...")
        return

    # 创建提取器
    extractor = PDFWordExtractor(pdf_path)

    # 提取文本
    text = extractor.extract_text()

    # 解析单词
    words = extractor.parse_words(text)

    if not words:
        print("❌ 没有提取到任何单词，请检查PDF格式")
        input("按回车键退出...")
        return

    # 显示示例
    print("\n" + "=" * 60)
    print("📝 提取结果预览（前5个单词）:")
    print("=" * 60)
    for i, word in enumerate(words[:5], 1):
        print(f"\n{i}. 单词: {word.word}")
        if word.phonetic:
            print(f"   音标: {word.phonetic}")
        if word.definition:
            print(f"   释义: {word.definition[:60]}...")
        if word.example_en:
            print(f"   例句: {word.example_en[:60]}...")

    # 选择导出字段
    print("\n" + "=" * 60)
    print("请选择要导出的字段（多选用逗号分隔，直接回车默认全选）:")
    print("1. 单词")
    print("2. 音标")
    print("3. 释义")
    print("4. 英文例句")
    print("5. 中文例句")

    choice = input("\n输入选择（例如: 1,2,3 或直接回车）: ").strip()

    field_map = {
        '1': 'word',
        '2': 'phonetic',
        '3': 'definition',
        '4': 'example_en',
        '5': 'example_cn',
    }

    if choice:
        selected = [field_map[c.strip()] for c in choice.split(',') if c.strip() in field_map]
    else:
        selected = list(field_map.values())

    # 选择导出格式
    print("\n" + "=" * 60)
    print("选择导出格式:")
    print("1. Excel (推荐，可直接导入应用)")
    print("2. TXT (纯文本)")
    print("3. 两者都要")

    format_choice = input("\n输入选择 (1/2/3): ").strip() or "1"

    # 生成输出文件名
    pdf_name = Path(pdf_path).stem
    output_dir = Path(pdf_path).parent

    if format_choice in ['1', '3']:
        excel_path = output_dir / f"{pdf_name}_单词表.xlsx"
        extractor.export_to_excel(excel_path, selected)

    if format_choice in ['2', '3']:
        txt_path = output_dir / f"{pdf_name}_单词表.txt"
        extractor.export_to_txt(txt_path)

    print("\n" + "=" * 60)
    print("🎉 完成！")
    print("=" * 60)
    input("\n按回车键退出...")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"\n❌ 发生错误: {e}")
        import traceback
        traceback.print_exc()
        input("\n按回车键退出...")
